"""
TrendLens — Taxonomy Extraction

Reads an uploaded report (PPTX or PDF), extracts its text, and uses Claude to
extract the trend taxonomy ENTIRELY FROM THE FILE (nothing from outside):

    Megatrend
      └─ Subtrend
           ├─ Products       (with a normalized search_term)
           ├─ Ingredients    (with a normalized search_term)
           ├─ Behaviours
           └─ Psychographics

Writes one `reports` row + many `taxonomy_nodes` rows. Prints REPORT_ID=<id>.

Usage:
    python -m src.extract_taxonomy "docs/Future of Protein Full Report 03.18.25.pptx"
"""

import argparse
import hashlib
import json
import logging
import shutil
import tempfile
import uuid
from datetime import datetime, timezone
from pathlib import Path

from pptx import Presentation
from pypdf import PdfReader

from src.common import bq_client, complete_json, dataset_ref, load_json_rows

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
logger = logging.getLogger(__name__)

MAX_TEXT_CHARS = 120_000  # cap text sent to the LLM

LEVELS_CHILD = ["products", "ingredients", "behaviours", "psychographics"]
LEVEL_SINGULAR = {
    "products": "product", "ingredients": "ingredient",
    "behaviours": "behaviour", "psychographics": "psychographic",
}


# ---------------------------------------------------------------------------
# Text extraction
# ---------------------------------------------------------------------------

def _extract_pptx(path: Path) -> str:
    # Copy to temp first — the original may be locked (e.g. open in PowerPoint / OneDrive).
    with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as tmp:
        shutil.copy(path, tmp.name)
        tmp_path = tmp.name
    parts: list[str] = []
    prs = Presentation(tmp_path)
    for slide_num, slide in enumerate(prs.slides, start=1):
        parts.append(f"[SLIDE {slide_num}]")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        parts.append(" | ".join(cells))
    Path(tmp_path).unlink(missing_ok=True)
    return "\n".join(parts)


def _extract_pdf(path: Path) -> str:
    # [PAGE n] markers let downstream synthesis cite the exact page (#page=N deep links).
    reader = PdfReader(str(path))
    return "\n".join(
        f"[PAGE {num}]\n" + (page.extract_text() or "")
        for num, page in enumerate(reader.pages, start=1)
    )


def extract_text(path: Path) -> tuple[str, str]:
    """Return (text, source_type)."""
    suffix = path.suffix.lower()
    if suffix == ".pptx":
        return _extract_pptx(path), "pptx"
    if suffix == ".pdf":
        return _extract_pdf(path), "pdf"
    raise ValueError(f"Unsupported file type: {suffix} (only .pptx and .pdf)")


# ---------------------------------------------------------------------------
# Claude extraction
# ---------------------------------------------------------------------------

EXTRACTION_PROMPT = """You are a trends analyst. Below is the full text of a market/innovation report. Extract its trend taxonomy STRICTLY FROM THIS DOCUMENT — do not add anything that is not supported by the text.

Structure to extract:
- Megatrends: the top-level forces/themes the report is organized around.
- For each megatrend, its Subtrends (the growth pathways / sub-themes under it).
- For each subtrend, four kinds of detail AS STATED OR CLEARLY IMPLIED in the report:
  - products: specific food/product types or formats mentioned
  - ingredients: specific ingredients/components mentioned
  - behaviours: what consumers DO — observable actions
  - psychographics: what consumers BELIEVE, VALUE, or how they SEE THEMSELVES — attitudes, values, mindsets, identities

For every product and ingredient, also give a `search_term`: a short, normalized phrase a normal person would type into Google for that thing (e.g. "pre-seasoned slow-cooked birria components" -> "birria"; "protein-fortified beverages" -> "protein water").

CRITICAL — behaviours vs. psychographics must be clean and distinct:
- A BEHAVIOUR is an observable action. Phrase it as a short present-tense verb phrase starting with a verb: "snacks between meals", "tracks macros with an app", "buys direct from the manufacturer". NOT a value, NOT a feeling, NOT a sentence fragment.
- A PSYCHOGRAPHIC is an attitude/value/identity. Phrase it as a short noun phrase: "values authenticity", "frugality", "health as identity", "convenience-seeking". NOT an action.
- DISTILL, do not quote. Turn the report's prose into a clean canonical label — do not lift mid-sentence clauses (bad: "more frugality trends and hacks emerge"; good: "frugality").
- If a single item is really both, put the action under behaviours and the value under psychographics.
- For each behaviour and psychographic, give a one-line `description` grounding it in what the report says.

## Report text
\"\"\"
{document_text}
\"\"\"

## Output
Return ONLY valid JSON, no markdown fencing:

{{
  "title": "the report's title or overall theme",
  "document_date": "the date or time period the report states ABOUT ITSELF, written exactly as it appears in the document (e.g. \\"2026/27\\", \\"March 2025\\", \\"Q1 2024\\"). Prefer the period the report covers/forecasts over the publication date; if only a publication date is stated, use that. Use null if the document does not state any date about itself. NEVER invent or infer a date that is not written in the text.",
  "megatrends": [
    {{
      "name": "Megatrend name",
      "description": "1 sentence from the report",
      "subtrends": [
        {{
          "name": "Subtrend name",
          "description": "1 sentence",
          "products": [{{"name": "as written", "search_term": "normalized"}}],
          "ingredients": [{{"name": "as written", "search_term": "normalized"}}],
          "behaviours": [{{"name": "verb phrase", "description": "1 line grounded in the report"}}],
          "psychographics": [{{"name": "value/identity noun phrase", "description": "1 line grounded in the report"}}]
        }}
      ]
    }}
  ]
}}

Rules:
- Only include items grounded in the document. Empty arrays are fine if the report doesn't mention that kind of detail for a subtrend.
- Keep names concise (behaviours/psychographics: at most ~6 words). Descriptions max 1 sentence.
- document_date must be copied verbatim from the document or null — never guessed.
- Return ONLY the JSON object.
"""


# Focused, non-destructive date extraction — used by the backfill script to populate
# document_date on already-extracted reports without regenerating their taxonomy nodes.
DATE_PROMPT = """Below is the text of a market/innovation report. Find the date or time period the report states ABOUT ITSELF — the period it covers/forecasts, or its publication date.

Return the date EXACTLY as written in the document (e.g. "2026/27", "March 2025", "Q1 2024"). Prefer the period the report covers/forecasts over the publication date. If the document states no date about itself, return null. NEVER invent or infer a date that is not written in the text.

## Report text
\"\"\"
{document_text}
\"\"\"

## Output
Return ONLY valid JSON, no markdown fencing:

{{"document_date": "<date as written, or null>"}}
"""


def extract_document_date(text: str) -> str | None:
    """Ask Claude for just the date the report states about itself (or None)."""
    result = complete_json(DATE_PROMPT.format(document_text=text[:MAX_TEXT_CHARS]))
    val = result.get("document_date") if isinstance(result, dict) else None
    return val.strip() if isinstance(val, str) and val.strip() else None


def extract_taxonomy(text: str) -> dict:
    text = text[:MAX_TEXT_CHARS]
    logger.info("Sending %d chars to Claude for taxonomy extraction...", len(text))
    return complete_json(EXTRACTION_PROMPT.format(document_text=text))


# ---------------------------------------------------------------------------
# Normalize + dedup behaviours / psychographics across the whole report (B)
# ---------------------------------------------------------------------------

NORMALIZE_PROMPT = """You are cleaning up a list of consumer {kind} extracted from a single trends report. Many items are near-duplicates phrased differently, and some are messy sentence fragments.

Produce a CANONICAL vocabulary: merge items that mean the same thing into one clean label, and rewrite messy items into a clean form.

A {kind} should be phrased as: {form}

## Raw items (JSON list)
{items}

## Output
Return ONLY valid JSON, no markdown fencing — a mapping from EVERY raw item (verbatim) to its canonical label. Items that mean the same thing must map to the SAME canonical label.

{{
  "mapping": {{
    "<raw item exactly as given>": "<clean canonical label>"
  }}
}}

Rules:
- Every raw item must appear as a key, exactly as written.
- Canonical labels are concise (at most ~6 words) and follow the form above.
- Merge aggressively when meaning matches; keep distinct when meaning differs.
- Return ONLY the JSON object.
"""

_NORMALIZE_FORM = {
    "behaviours": "a short present-tense verb phrase describing an observable action (e.g. \"snacks between meals\", \"buys direct from the manufacturer\").",
    "psychographics": "a short noun phrase describing a value, attitude, or identity (e.g. \"values authenticity\", \"frugality\", \"health as identity\").",
}


def _canonical_map(names: list[str], kind: str) -> dict[str, str]:
    """Ask Claude to merge near-duplicate names into a canonical label. Maps raw -> canonical."""
    uniq = sorted({n for n in names if n})
    if len(uniq) < 2:
        return {n: n for n in uniq}
    logger.info("Normalizing %d unique %s into canonical labels...", len(uniq), kind)
    result = complete_json(NORMALIZE_PROMPT.format(
        kind=kind, form=_NORMALIZE_FORM[kind], items=json.dumps(uniq, ensure_ascii=False)))
    mapping = result.get("mapping", {}) if isinstance(result, dict) else {}
    # Fall back to the original name for anything the model dropped.
    return {n: (mapping.get(n) or n).strip() for n in uniq}


def normalize_attributes(taxonomy: dict) -> dict:
    """Merge near-duplicate behaviours/psychographics across the report into canonical labels."""
    for plural in ("behaviours", "psychographics"):
        names: list[str] = []
        for mt in taxonomy.get("megatrends", []):
            for st in mt.get("subtrends", []):
                for item in st.get(plural, []) or []:
                    nm = (item.get("name") if isinstance(item, dict) else str(item)) or ""
                    if nm.strip():
                        names.append(nm.strip())
        if not names:
            continue
        cmap = _canonical_map(names, plural)
        for mt in taxonomy.get("megatrends", []):
            for st in mt.get("subtrends", []):
                cleaned, seen = [], set()
                for item in st.get(plural, []) or []:
                    if not isinstance(item, dict):
                        item = {"name": str(item)}
                    raw = (item.get("name") or "").strip()
                    if not raw:
                        continue
                    canon = cmap.get(raw, raw)
                    key = canon.lower()
                    if key in seen:  # dedup within this subtrend
                        continue
                    seen.add(key)
                    cleaned.append({"name": canon, "description": item.get("description")})
                st[plural] = cleaned
    return taxonomy


# ---------------------------------------------------------------------------
# Flatten to taxonomy_nodes rows
# ---------------------------------------------------------------------------

def build_node_rows(taxonomy: dict, report_id: str, now: str) -> list[dict]:
    rows: list[dict] = []

    def node(level, name, parent_id, megatrend, subtrend, description=None, search_term=None):
        nid = f"{level[:2]}_{uuid.uuid4().hex[:10]}"
        rows.append({
            "node_id": nid, "report_id": report_id, "parent_id": parent_id,
            "level": level, "name": name, "description": description,
            "search_term": search_term, "megatrend_name": megatrend,
            "subtrend_name": subtrend, "created_at": now,
        })
        return nid

    for mt in taxonomy.get("megatrends", []):
        mt_name = (mt.get("name") or "").strip()
        if not mt_name:
            continue
        mt_id = node("megatrend", mt_name, None, mt_name, None, mt.get("description"))
        for st in mt.get("subtrends", []):
            st_name = (st.get("name") or "").strip()
            if not st_name:
                continue
            st_id = node("subtrend", st_name, mt_id, mt_name, st_name, st.get("description"))
            for plural in LEVELS_CHILD:
                singular = LEVEL_SINGULAR[plural]
                for item in st.get(plural, []) or []:
                    iname = (item.get("name") or "").strip() if isinstance(item, dict) else str(item).strip()
                    if not iname:
                        continue
                    sterm = item.get("search_term") if isinstance(item, dict) else None
                    idesc = item.get("description") if isinstance(item, dict) else None
                    node(singular, iname, st_id, mt_name, st_name,
                         description=idesc, search_term=sterm)
    return rows


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def _existing_text_hash(client, project_id, dataset_id, report_id) -> str | None:
    """Return the stored text_hash for this report_id, or None if it doesn't exist yet."""
    from google.cloud import bigquery
    rows = list(client.query(
        f"SELECT text_hash FROM `{project_id}.{dataset_id}.reports` WHERE report_id = @r LIMIT 1",
        job_config=bigquery.QueryJobConfig(
            query_parameters=[bigquery.ScalarQueryParameter("r", "STRING", report_id)]),
    ).result())
    return rows[0].text_hash if rows else None


def run(file_path: str, report_id: str | None = None, filename: str | None = None) -> str:
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(file_path)

    text, source_type = extract_text(path)
    if not text.strip():
        raise ValueError("No text could be extracted from the file")
    display_name = filename or path.name
    logger.info("Extracted %d chars from %s (%s)", len(text), display_name, source_type)

    report_id = report_id or f"rep_{uuid.uuid4().hex[:12]}"
    text_hash = hashlib.sha256(text.encode("utf-8")).hexdigest()
    client = bq_client()
    project_id, dataset_id = dataset_ref()

    # Replace-in-place is keyed on filename (assigned upstream). Here we only decide whether
    # the content actually changed: if this report already has the same text, skip the
    # expensive re-extraction entirely (no LLM tokens). Otherwise (new file or a revision)
    # we extract and overwrite the report's rows below.
    if _existing_text_hash(client, project_id, dataset_id, report_id) == text_hash:
        logger.info("Content unchanged for %s (text_hash %s…) — skipping re-extraction", report_id, text_hash[:12])
        print(f"REPORT_ID={report_id}")
        return report_id

    taxonomy = extract_taxonomy(text)
    taxonomy = normalize_attributes(taxonomy)

    now = datetime.now(timezone.utc).isoformat()
    node_rows = build_node_rows(taxonomy, report_id, now)
    num_megatrends = sum(1 for r in node_rows if r["level"] == "megatrend")

    # Idempotent: clear any prior rows for this report_id (handles re-runs / placeholders).
    from google.cloud import bigquery
    for tbl in ("reports", "taxonomy_nodes"):
        client.query(
            f"DELETE FROM `{project_id}.{dataset_id}.{tbl}` WHERE report_id = @rid",
            job_config=bigquery.QueryJobConfig(
                query_parameters=[bigquery.ScalarQueryParameter("rid", "STRING", report_id)]),
        ).result()

    report_row = {
        "report_id": report_id,
        "filename": display_name,
        "title": taxonomy.get("title"),
        "document_date": (taxonomy.get("document_date") or None) if isinstance(taxonomy.get("document_date"), str) and taxonomy.get("document_date").strip() else None,
        "source_type": source_type,
        "num_megatrends": num_megatrends,
        "num_nodes": len(node_rows),
        "raw_text_chars": len(text),
        "status": "extracted",
        "text_hash": text_hash,
        "uploaded_at": now,
    }
    load_json_rows(client, f"{project_id}.{dataset_id}.reports", [report_row])
    load_json_rows(client, f"{project_id}.{dataset_id}.taxonomy_nodes", node_rows)

    logger.info("Extracted %d megatrends, %d total nodes -> report %s",
                num_megatrends, len(node_rows), report_id)
    for r in node_rows:
        if r["level"] == "megatrend":
            subs = [n["name"] for n in node_rows if n["parent_id"] == r["node_id"]]
            logger.info("  %s  (%d subtrends)", r["name"], len(subs))
    print(f"REPORT_ID={report_id}")
    return report_id


def main() -> None:
    parser = argparse.ArgumentParser(description="TrendLens — extract taxonomy from a report file")
    parser.add_argument("file", help="Path to a .pptx or .pdf report")
    parser.add_argument("--report-id", default=None, help="Use a specific report id (replace in place)")
    parser.add_argument("--filename", default=None, help="Original filename to record (else the file's name)")
    args = parser.parse_args()
    run(args.file, report_id=args.report_id, filename=args.filename)


if __name__ == "__main__":
    main()
